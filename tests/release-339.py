"""Focused plans, confirmations, feed isolation and scoped conversations; mocked providers."""
import os,uuid,unittest,shutil
from pathlib import Path
from unittest.mock import patch,Mock
root=Path(__file__).resolve().parents[1];temp=root/('zar-339-'+uuid.uuid4().hex);temp.mkdir()
env=patch.dict(os.environ,{'ZAR_DATA_DIR':str(temp),'RAILWAY_ENVIRONMENT_ID':''});env.start()
with patch('threading.Thread.start'):
 from app import main,holdings,semantic_tasks as tasks,semantic_planner as planner,confirmations,memory
 from app.user_scope import set_current_user
 from app.stonks_realtime import StonksRealtimeBus

class Release339(unittest.TestCase):
 def setUp(self):self.scope='fixture_'+uuid.uuid4().hex;set_current_user(self.scope);self.cfg=patch('app.config.load',return_value={});self.cfg.start()
 def tearDown(self):self.cfg.stop()
 def test_multistep_dependencies(self):
  task=tasks.create(self.scope,'Investiga ayudas para jóvenes emprendedores, después haz un informe y envíaselo por email a Belloso.')
  kinds=[s['kind'] for s in task['steps']]
  self.assertEqual(kinds,['RESEARCH','CREATE_REPORT','RESOLVE_CONTACT','DRAFT_EMAIL','ATTACH_ARTIFACT','SEND_EMAIL'])
  for s in task['steps']:self.assertTrue(all(int(dep)<s['order'] for dep in s['dependencies']))
  self.assertEqual(task['entities']['recipient'],'Belloso');self.assertNotIn('READ_MAIL',kinds)
 def test_disambiguation(self):
  for message,kind in [('mira mi correo','READ_MAIL'),('manda esto por correo','SEND_EMAIL'),('el correo de Belloso','RESOLVE_CONTACT'),('adjunta el informe al correo','ATTACH_ARTIFACT'),('hazme un correo explicándole esto','DRAFT_EMAIL'),('mete esas cifras en una hoja','CREATE_REPORT'),('usa el documento que hicimos ayer','ATTACH_ARTIFACT')]:
   self.assertIn(kind,planner.build(self.scope,message,tasks.plan,False)['tools'])
 def test_reference_waits_without_evidence(self):
  task=tasks.create(self.scope,'haz un informe con lo anterior');result=tasks.run(self.scope,task['id']);self.assertEqual(result['status'],'WAITING');self.assertNotIn('CREATE_REPORT',result['outputs']);self.assertEqual(result['steps'],result['subtasks'])
 def test_followup_uses_actual_task(self):
  first=tasks.create(self.scope,'Investiga una fuente');first.update(status='DONE',outputs={'RESEARCH':{'text':'Evidence','sources':[{'url':'https://example.test'}]}});tasks.save(self.scope,first)
  follow=tasks.create(self.scope,'haz un informe con lo anterior');self.assertEqual(follow['references'][0]['task_id'],first['id']);self.assertEqual(planner.resolve_reference(self.scope,follow)['research']['text'],'Evidence')
 def test_model_cannot_invent_external_action_or_reverse_dependencies(self):
  self.cfg.stop()
  for steps in [[{'kind':'CREATE_REPORT'},{'kind':'RESEARCH'}],[{'kind':'RESEARCH'},{'kind':'RESOLVE_CONTACT'},{'kind':'DRAFT_EMAIL'},{'kind':'SEND_EMAIL'}]]:
   import json
   with patch('app.config.load',return_value={'api':{'api_key':'fixture'}}),patch('app.agent.api_text',return_value=json.dumps({'steps':steps})):
    p=planner.build(self.scope,'Investiga una fuente y haz un informe',tasks.plan);self.assertEqual(p['planning_engine'],'VERB_OBJECT_CONTEXT');self.assertEqual(p['tools'],['RESEARCH','CREATE_REPORT'])
  self.cfg.start()
 def test_high_requires_separate_decisions_and_idempotent_claim(self):
  payload={'kind':'email','to':'fixture@example.test'};row=confirmations.prepare(self.scope,'Email sensible',payload,'HIGH')
  self.assertEqual(row['confirmation_id'],confirmations.prepare(self.scope,'Email sensible',payload,'HIGH')['confirmation_id'])
  first=confirmations.decide(self.scope,row['confirmation_id'],payload,'confirm',True,'first');self.assertEqual(first['state'],'WAITING_SECOND');self.assertIsNone(first['job_id'])
  second=confirmations.decide(self.scope,row['confirmation_id'],payload,'confirm',True,'second');self.assertEqual(second['state'],'RUNNING')
  repeat=confirmations.decide(self.scope,row['confirmation_id'],payload,'confirm',True,'duplicate');self.assertEqual(repeat['job_id'],'second')
  confirmations.complete(self.scope,row['confirmation_id']);self.assertEqual(holdings.read(self.scope)['confirmations'][0]['state'],'EXECUTED')
 def test_payload_binding_and_owner(self):
  row=confirmations.prepare(self.scope,'Action',{'id':1})
  with self.assertRaises(ValueError):confirmations.decide(self.scope,row['confirmation_id'],{'id':2},'confirm')
  with self.assertRaises(ValueError):confirmations.decide('other',row['confirmation_id'],{'id':1},'confirm')
 def test_action_event_not_user_message(self):
  with main.app.test_request_context('/'):
   token=main._action_event.set('ACTION_CONFIRMED: Pablo confirmó fixture')
   try:
    with patch.object(main,'remember'),patch.object(main,'add_message') as messages,patch.object(main,'add_conversation_message'):
     main._remember_turn('user','__ZAR_CONFIRM__');self.assertEqual(messages.call_args.args,('system','ACTION_CONFIRMED: Pablo confirmó fixture'))
   finally:main._action_event.reset(token)
 def test_bus_private_events_and_symbol_snapshot(self):
  bus=StonksRealtimeBus();one,q1=bus.subscribe('one');two,q2=bus.subscribe('two');bus.publish('one','PORTFOLIO_UPDATE',{'cash':1});self.assertEqual(q1.get_nowait()['data']['cash'],1);self.assertTrue(q2.empty())
  for symbol in ['BTC/USD','ETH/USD']:bus.publish('MARKET_PUBLIC','MARKET_TICK',{'symbol':symbol,'price':100})
  self.assertEqual(len(bus.snapshot('two')),2);self.assertTrue(all(x['zero_tokens'] for x in bus.snapshot('one').values()));bus.unsubscribe(one);bus.unsubscribe(two)
 def test_actual_ohlcv_provider_contract(self):
  with main.app.test_request_context('/api/stonks/chart/btc?timeframe=5Min'),patch.object(main,'_alpaca_market_request',return_value={'bars':{'BTC/USD':[{'t':'2026-10-07T10:00:00Z','o':1,'h':2,'l':1,'c':2,'v':3}]}}) as request:
   row=main.stonks_btc_chart_api().get_json();self.assertEqual(row['symbol'],'BTC/USD');self.assertEqual(row['bars'][0]['v'],3);self.assertTrue(row['zero_tokens']);self.assertEqual(request.call_args.kwargs['params']['timeframe'],'5Min')
 def test_conversations_scoped_metadata(self):
  memory.add_conversation_message('user','Fixture conversación');memory.add_conversation_message('assistant','Snippet fixture');row=memory.archive_current_conversation();listed=memory.list_conversations()[0];self.assertEqual(listed['message_count'],2);self.assertEqual(listed['snippet'],'Snippet fixture')
  memory.update_conversation_archive(row['id'],title='Renamed',starred=True);self.assertTrue(memory.list_conversations()[0]['starred']);set_current_user('other');self.assertIsNone(memory.get_conversation_archive(row['id']));set_current_user(self.scope);self.assertTrue(memory.delete_conversation_archive(row['id']))
 def test_charts_and_professional_workbook(self):
  import io
  from app.chart_agent import render
  from app.artifact_engine import create
  from app.file_store import FileStore
  from openpyxl import load_workbook
  for kind in ['bar','line','area','pie','donut','scatter']:self.assertTrue(render({'type':kind,'values':[2,3,4],'labels':['A','B','C']}).startswith(b'\x89PNG'))
  with self.assertRaises(ValueError):render({'type':'line','values':[float('nan')]})
  artifact=create('Workbook fixture','Nombre\tImporte\nA\t12\nB\t18',kind='xlsx',template='FINANCIAL_REPORT',charts=[{'type':'bar','values':[12,18],'labels':['A','B']}]);wb=load_workbook(io.BytesIO(FileStore().read(artifact['artifact_id'])))
  self.assertIn('Dashboard',wb.sheetnames);self.assertIn('Gráfico 1',wb.sheetnames);self.assertEqual(wb['Análisis']['B2'].value,12);self.assertEqual(wb['Dashboard']['B2'].data_type,'f');self.assertEqual(wb['Análisis'].freeze_panes,'A2');self.assertTrue(wb['Dashboard'].data_validations.count)
 def test_workspace_verified_writes_and_no_duplicate(self):
  from app import workspace_verification as probe
  docs=Mock();docs.documents.return_value.get.return_value.execute.return_value={'body':{'content':'ZAR DOCS OK'}}
  sheets=Mock();sheets.spreadsheets.return_value.get.return_value.execute.return_value={'sheets':[{'properties':{'sheetId':42}}]};sheets.spreadsheets.return_value.values.return_value.get.return_value.execute.return_value={'values':[['ZAR','SHEETS','OK'],[2]]}
  slides=Mock();slides.presentations.return_value.get.return_value.execute.side_effect=[{'slides':[]},{'slides':[{'objectId':'zar_workspace_test','content':'SLIDES OK'}]}]
  drive=Mock();drive.files.return_value.list.return_value.execute.return_value={'files':[{'id':'fixture-folder'}]}
  with patch.object(probe.identity_center,'require_mail',return_value='zaragente031@gmail.com'),patch.object(probe.w,'docs_create',return_value={'documentId':'fixture-doc','url':'https://docs.google.com/document/d/fixture-doc'}) as create,patch.object(probe.w,'sheets_create',return_value={'spreadsheetId':'fixture-sheet','url':'https://docs.google.com/spreadsheets/d/fixture-sheet'}),patch.object(probe.w,'slides_create',return_value={'presentationId':'fixture-slides','url':'https://docs.google.com/presentation/d/fixture-slides'}),patch.object(probe.w,'docs_service',return_value=docs),patch.object(probe.w,'sheets_service',return_value=sheets),patch.object(probe.w,'slides_service',return_value=slides),patch.object(probe.w,'drive_service',return_value=drive):
   rows=probe.verify(self.scope);self.assertTrue(all(r['status']=='CONNECTED' for r in rows.values()));probe.verify(self.scope);create.assert_called_once();self.assertEqual(holdings.read(self.scope)['identity_center']['capabilities']['DOCS']['document_id'],'fixture-doc');self.assertEqual(sheets.spreadsheets.return_value.values.return_value.update.call_args.kwargs['valueInputOption'],'USER_ENTERED')

if __name__=='__main__':
 try:result=unittest.TextTestRunner().run(unittest.defaultTestLoader.loadTestsFromTestCase(Release339))
 finally:env.stop();shutil.rmtree(temp,ignore_errors=True)
 if not result.wasSuccessful():raise SystemExit(1)
